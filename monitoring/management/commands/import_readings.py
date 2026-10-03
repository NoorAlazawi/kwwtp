"""
أمر إدارة لاستيراد القراءات اليومية من ملف الإكسل إلى قاعدة البيانات.

الاستخدام:
    python manage.py import_readings --file data/raw_data.xlsx --clarifier 2
    python manage.py import_readings --file data/raw_data.xlsx --clarifier 2 --dry-run

مقاوم لإزاحة الأعمدة: يتعرّف على كل عمود من اسمه في صفّ الرؤوس (لا من ترتيبه)،
ويتجاهل تلقائياً الأعمدة المحسوبة (Qr، Qw، SVI) لأن النظام يحسبها بنفسه.
"""
import re
from decimal import Decimal, InvalidOperation
from datetime import date as _date, datetime
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from monitoring.models import Plant, Clarifier, DailyReading


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9%]", "", str(s).lower())


FIELD_RULES = [
    ("flow_q",         lambda h: "influentflow" in h or ("flow" in h and "q" in h and "return" not in h and "waste" not in h)),
    ("ras_percent",    lambda h: h in ("%ras", "ras%") or ("ras" in h and "%" in h)),
    ("was_percent",    lambda h: h in ("%was", "was%") or ("was" in h and "%" in h)),
    ("mlss",           lambda h: "mlss" in h or "mixedliquor" in h),
    ("sv30",           lambda h: "sv30" in h or ("sv" in h and "30" in h)),
    ("do",             lambda h: "dissolvedoxygen" in h or h == "do" or h.startswith("do")),
    ("temperature",    lambda h: "temperature" in h or h in ("t", "temp")),
    ("sludge_blanket", lambda h: "sludgeblanket" in h or "blanket" in h),
    ("tss_in",         lambda h: "tss" in h and "in" in h and "out" not in h),
    ("tss_out",        lambda h: "tss" in h and "out" in h),
    ("cod_in",         lambda h: "cod" in h and "in" in h and "out" not in h),
    ("cod_out",        lambda h: "cod" in h and "out" in h),
    ("bod_in",         lambda h: "bod" in h and "in" in h and "out" not in h),
    ("bod_out",        lambda h: "bod" in h and "out" in h),
    ("ph",             lambda h: h == "ph"),
]

COMPUTED_HINTS = ("returnactivatedsludgerate", "wasteactivatedsludgerate",
                  "sludgevolumeindex", "svi", "qrm3", "qwm3")
DATE_HINTS = ("data", "date", "day")
CLARIFIER_HINTS = ("clarifier", "basin", "tank", "st-")


class Command(BaseCommand):
    help = "استيراد القراءات اليومية من ملف Excel — يتعرّف على الأعمدة بأسمائها"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="مسار ملف الإكسل")
        parser.add_argument("--sheet", default="Raw data", help="اسم الورقة (افتراضي: Raw data)")
        parser.add_argument("--clarifier", type=int, default=None,
                            help="رقم الحوض لكل الصفوف (يُتجاهل إن وُجد عمود حوض في الملف)")
        parser.add_argument("--plant", default="Karbala WWTP", help="اسم المحطة")
        parser.add_argument("--dry-run", action="store_true", help="فحص دون حفظ")
        parser.add_argument("--update", action="store_true", help="تحديث الموجود بنفس التاريخ")

    def _map_headers(self, header_row):
        field_idx, used = {}, set()
        date_idx = clar_idx = None
        norm = [_norm(h) if h is not None else "" for h in header_row]
        for i, h in enumerate(norm):
            if not h:
                continue
            if date_idx is None and any(k in h for k in DATE_HINTS):
                date_idx = i; continue
            if clar_idx is None and any(k in h for k in CLARIFIER_HINTS) and "%" not in h:
                clar_idx = i; continue
        for field, rule in FIELD_RULES:
            for i, h in enumerate(norm):
                if i in used or not h or i in (date_idx, clar_idx):
                    continue
                if any(k in h for k in COMPUTED_HINTS):
                    continue
                if rule(h):
                    field_idx[field] = i; used.add(i); break
        return field_idx, date_idx, clar_idx

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("مكتبة openpyxl غير مثبّتة. نفّذ: pip install openpyxl")
        try:
            wb = openpyxl.load_workbook(opts["file"], data_only=True)
        except FileNotFoundError:
            raise CommandError(f"الملف غير موجود: {opts['file']}")
        if opts["sheet"] not in wb.sheetnames:
            raise CommandError(f"الورقة '{opts['sheet']}' غير موجودة. الأوراق: {wb.sheetnames}")
        ws = wb[opts["sheet"]]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise CommandError("الورقة فارغة.")
        header, data_rows = rows[0], rows[1:]
        field_idx, date_idx, clar_idx = self._map_headers(header)

        missing = [f for f, _ in FIELD_RULES if f not in field_idx]
        if missing:
            self.stdout.write(self.style.WARNING(
                "تعذّر التعرّف على أعمدة: " + ", ".join(missing) +
                "\nالرؤوس الموجودة: " + " | ".join(str(h) for h in header if h)))
            raise CommandError("تأكّد من صفّ الرؤوس في الملف.")
        if date_idx is None:
            raise CommandError("لم يُعثر على عمود التاريخ.")

        self.stdout.write(self.style.SUCCESS(
            f"تعرّفت على {len(field_idx)} عمود إدخال · عمود التاريخ #{date_idx+1}" +
            (f" · عمود الحوض #{clar_idx+1}" if clar_idx is not None else " · بلا عمود حوض")))

        plant, _ = Plant.objects.get_or_create(name=opts["plant"], defaults={"num_clarifiers": 8})
        clar_cache = {}
        def get_clar(num):
            num = int(num)
            if num not in clar_cache:
                c, cr = Clarifier.objects.get_or_create(
                    plant=plant, number=num, defaults={"diameter_m": 32, "side_water_depth": 4})
                if cr:
                    self.stdout.write(self.style.WARNING(f"أُنشئ حوض ST-{num}."))
                clar_cache[num] = c
            return clar_cache[num]

        default_clar = None
        if clar_idx is None:
            if opts["clarifier"] is None:
                raise CommandError("لا يوجد عمود حوض في الملف — حدّد --clarifier N.")
            default_clar = get_clar(opts["clarifier"])

        created_n = updated_n = skipped_n = error_n = 0
        with transaction.atomic():
            for i, row in enumerate(data_rows, start=2):
                date_val = row[date_idx] if date_idx < len(row) else None
                q_raw = row[field_idx["flow_q"]] if field_idx["flow_q"] < len(row) else None
                if date_val is None or q_raw is None:
                    continue
                if isinstance(date_val, (datetime, _date)):
                    the_date = date_val.date() if isinstance(date_val, datetime) else date_val
                else:
                    try:
                        the_date = datetime.fromisoformat(str(date_val)).date()
                    except ValueError:
                        error_n += 1
                        self.stderr.write(self.style.ERROR(f"صف {i}: تاريخ غير صالح '{date_val}'"))
                        continue
                clar = default_clar or get_clar(row[clar_idx])
                try:
                    values = {}
                    for field, idx in field_idx.items():
                        raw = row[idx] if idx < len(row) else None
                        if raw is None:
                            raise ValueError(f"قيمة فارغة للحقل {field}")
                        values[field] = Decimal(str(round(float(raw), 3)))
                except (ValueError, InvalidOperation) as e:
                    error_n += 1
                    self.stderr.write(self.style.ERROR(f"صف {i} ({the_date}): {e}"))
                    continue
                exists = DailyReading.objects.filter(clarifier=clar, date=the_date).exists()
                if exists and not opts["update"]:
                    skipped_n += 1; continue
                if opts["dry_run"]:
                    created_n += 0 if exists else 1
                    updated_n += 1 if exists else 0
                    continue
                _, wc = DailyReading.objects.update_or_create(
                    clarifier=clar, date=the_date, defaults=values)
                created_n += 1 if wc else 0
                updated_n += 0 if wc else 1
            if opts["dry_run"]:
                transaction.set_rollback(True)

        mode = "[فحص فقط] " if opts["dry_run"] else ""
        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS(
            f"{mode}تمّ: {created_n} جديدة · {updated_n} محدّثة · "
            f"{skipped_n} متخطّاة · {error_n} أخطاء"))
        self.stdout.write(f"إجمالي صفوف البيانات: {len(data_rows)}")
